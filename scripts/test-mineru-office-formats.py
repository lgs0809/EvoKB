"""Actual MinerU3.4.5 XLSX/PPTX parsing; original files/gold native results retained, no model claim."""
import argparse,hashlib,importlib.metadata,json,sys,time,traceback
from datetime import datetime,timezone
from pathlib import Path
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
R=Path(__file__).resolve().parents[1];p=argparse.ArgumentParser();p.add_argument('--site',type=Path,default=R/'.tmp/mineru-office/site');p.add_argument('--output',type=Path,required=True);a=p.parse_args()
if a.output.exists():raise SystemExit('Fresh evidence directory required')
a.output.mkdir(parents=True);sys.path.insert(0,str(a.site.resolve()));fixture=a.output/'formats';fixture.mkdir();marker=fixture/'marker.png';Image.new('RGB',(150,100),(24,90,160)).save(marker)
rows=[['合同编号','金额（万元）','质保（月）','付款条件'],['HT-900','900.00','24','合法发票且验收后60日'],['HT-902','120.50','12','合法发票且验收后30日']]
w=Workbook();s=w.active;s.title='付款计划';s.merge_cells('A1:D1');s['A1']='合成验收；金额单位万元，非实际业务数据'
for row in rows:s.append(row)
s.append(['表注','缺少合法发票不得付款','','']);s.merge_cells('B5:D5');xlsx=fixture/'payment.xlsx';w.save(xlsx)
ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[6]);title=slide.shapes.add_textbox(Inches(.4),Inches(.2),Inches(8),Inches(.6));title.text='付款条件（合成验收）'
text=slide.shapes.add_textbox(Inches(.4),Inches(1),Inches(8),Inches(.8));text.text='合同HT-900金额900万元，质保24个月。合法发票且完成验收后60日内付款；缺少合法发票不得付款。'
table=slide.shapes.add_table(3,4,Inches(.4),Inches(2.2),Inches(8.5),Inches(2)).table
for i,row in enumerate(rows):
 for j,value in enumerate(row):table.cell(i,j).text=value
slide.shapes.add_picture(str(marker),Inches(.4),Inches(4.7),width=Inches(2),height=Inches(1.2));pptx=fixture/'payment.pptx';ppt.save(pptx)
gold=dict(rows=rows,criticalConditions=['合法发票','验收','60日'],boundary='Native XLSX/PPTX element order is not PDF page geometry. Formula recalculation, embedded charts, multi-slide ordering are not claimed.')
(a.output/'gold.json').write_text(json.dumps(gold,ensure_ascii=False,indent=2));report=dict(status='FAIL',startedAt=datetime.now(timezone.utc).isoformat(),checks={},results={},gold=gold,sourceSha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest());start=time.monotonic()
try:
 from mineru.version import __version__
 from mineru.backend.office.xlsx_analyze import office_xlsx_analyze
 from mineru.backend.office.pptx_analyze import office_pptx_analyze
 from mineru.backend.office.office_middle_json_mkcontent import union_make
 from mineru.utils.enum_class import MakeMode
 from bs4 import BeautifulSoup
 for name,path,analyze in [('xlsx',xlsx,office_xlsx_analyze),('pptx',pptx,office_pptx_analyze)]:
  folder=a.output/name;folder.mkdir()
  class Writer:
   def write(self,key,data):target=folder/'images'/key;target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(data)
  middle,model=analyze(path.read_bytes(),image_writer=Writer());content=union_make(middle['pdf_info'],MakeMode.CONTENT_LIST,'images');markdown=union_make(middle['pdf_info'],MakeMode.MM_MD,'images')
  for key,value in [('middle.json',middle),('content-list-v1.json',content),('native-model-output.json',model)]: (folder/key).write_text(json.dumps(value,ensure_ascii=False,indent=2,default=lambda x:x.model_dump() if hasattr(x,'model_dump') else str(x)))
  (folder/'document.md').write_text(markdown);tables=[BeautifulSoup(x.get('table_body',''),'html.parser') for x in content if x.get('type')=='table'];whole=' '.join(t.get_text(' ',strip=True) for t in tables)+' '+' '.join(x.get('text','') for x in content)
  report['checks'][name+'ActualMinerU345']=__version__=='3.4.5';report['checks'][name+'AllRowsAmountsAndUnits']=all(value in whole for value in ['HT-900','HT-902','900.00','120.50','万元','质保']);report['checks'][name+'IndependentConditionsPreserved']=all(value in whole for value in ['合法发票','验收','60日','30日']);report['checks'][name+'NativeSourceStructureRetained']=bool(content) and bool(middle.get('pdf_info')) and bool(tables)
  if name=='xlsx':report['checks']['xlsxMergedCellSpansRetained']=any(t.select_one('[colspan="4"]') for t in tables)
  else:report['checks']['pptxImageBytesPersisted']=any(x.is_file() for x in (folder/'images').rglob('*')) if (folder/'images').exists() else False
  report['results'][name]=dict(originalSha256=hashlib.sha256(path.read_bytes()).hexdigest(),blocks=len(content),nativePageIndices=[x.get('page_idx') for x in content],boundary=gold['boundary'])
 report['status']='PASS_ACTUAL_MINERU_OFFICE_FORMATS' if all(report['checks'].values()) else 'FAIL_ACTUAL_MINERU_OFFICE_FORMATS'
except Exception as error:(a.output/'failure.log').write_text(traceback.format_exc());report['error']=str(error)
report['elapsedSeconds']=round(time.monotonic()-start,3);report['packages']={x:dict(version=importlib.metadata.version(x),distributionPath=str(importlib.metadata.distribution(x)._path)) for x in ['mineru','python-pptx','openpyxl','pillow']};report['artifactHashes']={str(x.relative_to(a.output)):hashlib.sha256(x.read_bytes()).hexdigest() for x in a.output.rglob('*') if x.is_file()};(a.output/'result.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(dict(status=report['status'],checks=report['checks'],error=report.get('error')),ensure_ascii=False));raise SystemExit(0 if report['status'].startswith('PASS') else 1)
