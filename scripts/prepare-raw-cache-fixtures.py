"""Repeatable related DOCX versions for real backend/browser cache acceptance; no DB writes."""
import argparse,hashlib,json,uuid
from datetime import datetime,timezone
from pathlib import Path
from docx import Document
p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);a=p.parse_args()
manifest=a.output/'manifest.json'
if manifest.exists():
 current=json.loads(manifest.read_text())
 for record in current['versions']:
  if hashlib.sha256(Path(record['path']).read_bytes()).hexdigest()!=record['sha256']:raise SystemExit('Existing fixture changed; no overwrite performed')
 print(json.dumps(current,ensure_ascii=False));raise SystemExit(0)
a.output.mkdir(parents=True,exist_ok=True);run=uuid.uuid4().hex[:10];name='raw-cache-acceptance-'+run+'.docx';records=[]
for version,days,title in [(1,60,'尽调合同（合成验收）'),(2,30,'尽调合同（合成验收）'),(3,30,'尽调合同（标题修订）')]:
 doc=Document();doc.add_heading(title,level=1)
 for heading,text in [('第一条 金额与质保','合同HT-900：金额900万元，质保24个月。此条在三个版本中正文完全相同。'),('第二条 付款先决条件',f'未取得合法发票前付款义务不发生；取得合法发票且完成验收后{days}日内付款。'),('第三条 权限与重复请求','本资料为合成测试数据，所有下载与引用仍按当前用户权限校验；相同上传不应产生新的活动版本。')]:doc.add_heading(heading,level=2);doc.add_paragraph(text)
 folder=a.output/f'v{version}';folder.mkdir();path=folder/name;doc.save(path)
 records.append(dict(version=version,path=str(path.resolve()),filename=name,sha256=hashlib.sha256(path.read_bytes()).hexdigest(),days=days,title=title,expectedEncodingChange='all new inputs' if version==1 else 'only changed second clause; first/third exact inputs reusable' if version==2 else 'all heading-prefixed inputs change; body-only first/third hashes stay the same'))
result=dict(createdAt=datetime.now(timezone.utc).isoformat(),logicalName=name.rsplit('.',1)[0],versions=records,execution='Upload V1 normally, wait for normal ingestion ACTIVE/SUCCEEDED; upload same V1 once more to validate idempotency; then V2, wait; then V3. Keep source identity/logical name consistent. Do not upload all revisions concurrently or write success states to DB.',boundary='Fixtures only. Not a model/cache/browser or actual ingestion pass.')
manifest.write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False))
