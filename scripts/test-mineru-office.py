"""Actual MinerU 3.4.5 Office parser, independent gold facts; no PDF/vision-model claim."""
import argparse,hashlib,json,sys,time
from datetime import datetime,timezone
from pathlib import Path
import importlib.metadata
from docx import Document
from docx.shared import Inches
from PIL import Image

R=Path(__file__).resolve().parents[1]
p=argparse.ArgumentParser();p.add_argument('--site',type=Path,default=R/'.tmp/mineru-office/site');p.add_argument('--output',type=Path,required=True);a=p.parse_args()
if a.output.exists(): raise SystemExit('Fresh output required; failures are retained')
a.output.mkdir(parents=True)
sys.path.insert(0,str(a.site.resolve()))
fixture=a.output/'invoice-acceptance.docx';image=a.output/'synthetic-marker.png'
Image.new('RGB',(12,8),(40,90,170)).save(image)
doc=Document();doc.add_heading('尽调付款测试（合成数据）',level=1)
doc.add_paragraph('合同HT-900：金额900万元，质保24个月。未取得合法发票前付款义务不发生；取得合法发票且完成验收后60日内付款。')
doc.add_heading('第二条 回款计划',level=2)
table=doc.add_table(rows=5,cols=3);table.style='Table Grid'
table.cell(0,0).text='合同编号';table.cell(0,1).merge(table.cell(0,2)).text='付款计划（单位：万元）'
table.cell(0,0).merge(table.cell(1,0));table.cell(1,1).text='金额';table.cell(1,2).text='条件'
for index,(key,value,condition) in enumerate([('HT-900','900','发票且验收后60日'),('HT-901','0','暂缓付款'),('HT-902','120.50','质保24个月')],2):
 for col,text in enumerate((key,value,condition)):table.cell(index,col).text=text
doc.add_paragraph('表注：金额单位为万元；缺少合法发票时不得付款。')
doc.add_picture(str(image),width=Inches(.25));doc.save(fixture)
gold=dict(model_not_required=True,document_key='HT-900',amount='900',unit='万元',warranty='24个月',delay='60日',condition='合法发票且完成验收',table_keys=['HT-900','HT-901','HT-902'],rowspan=2,colspan=2)
(a.output/'gold.json').write_text(json.dumps(gold,ensure_ascii=False,indent=2))
started=time.monotonic();report=dict(createdAt=datetime.now(timezone.utc).isoformat(),status='FAIL',inputSha256=hashlib.sha256(fixture.read_bytes()).hexdigest(),gold=gold)
try:
 from mineru.version import __version__
 from mineru.backend.office.docx_analyze import office_docx_analyze
 from mineru.backend.office.office_middle_json_mkcontent import union_make
 from mineru.utils.enum_class import MakeMode
 class Writer:
  def write(self,path,data):
   target=a.output/'images'/path;target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(data)
 middle,model=office_docx_analyze(fixture.read_bytes(),image_writer=Writer())
 content=union_make(middle['pdf_info'],MakeMode.CONTENT_LIST,'images')
 markdown=union_make(middle['pdf_info'],MakeMode.MM_MD,'images')
 for name,value in [('middle.json',middle),('content-list-v1.json',content),('native-office-elements.json',model)]:
  (a.output/name).write_text(json.dumps(value,ensure_ascii=False,indent=2,default=lambda x:x.model_dump() if hasattr(x,'model_dump') else str(x)))
 (a.output/'document.md').write_text(markdown)
 from bs4 import BeautifulSoup
 tables=[BeautifulSoup(b.get('table_body',''),'html.parser') for b in content if b.get('type')=='table']
 body='\n'.join(b.get('text','') for b in content)
 tabletext='\n'.join(t.get_text(' ',strip=True) for t in tables)
 checks=dict(actualMinerU345=__version__=='3.4.5',headingHierarchy=any(b.get('text_level')==1 and '尽调付款测试' in b.get('text','') for b in content) and any(b.get('text_level')==2 for b in content),criticalTextFacts=all(x in body for x in ['900万元','24个月','60日','未取得合法发票前付款义务不发生','合法发票且完成验收']),allBusinessRows=all(x in tabletext for x in gold['table_keys']) and '120.50' in tabletext,mergedCellSpansRetained=any(t.select_one('[rowspan="2"]') and t.select_one('[colspan="2"]') for t in tables),tableUnitAndCondition='单位：万元' in tabletext and '发票且验收后60日' in tabletext,imageActuallyPersisted=any((a.output/'images').rglob('*')) if (a.output/'images').exists() else False)
 report.update(status='PASS_ACTUAL_MINERU_OFFICE' if all(checks.values()) else 'FAIL_ACTUAL_MINERU_OFFICE',minerUVersion=__version__,requestedBackend='hybrid-engine',effectiveBackend='native-office',elapsedSeconds=time.monotonic()-started,checks=checks,contentBlocks=len(content),imageFiles=[str(x.relative_to(a.output)) for x in (a.output/'images').rglob('*') if x.is_file()] if (a.output/'images').exists() else [],boundary='Real native MinerU3.4.5 DOCX parser. Office element order/page_idx is NOT validated PDF page geometry. No scan/PDF hybrid model inference or retrieval embedding performed.')
except Exception as error:
 import traceback
 (a.output/'failure.log').write_text(traceback.format_exc());report.update(error=str(error),elapsedSeconds=time.monotonic()-started)
report['sourceSha256']=hashlib.sha256(Path(__file__).read_bytes()).hexdigest();report['artifactHashes']={str(x.relative_to(a.output)):hashlib.sha256(x.read_bytes()).hexdigest() for x in a.output.rglob('*') if x.is_file()}
(a.output/'result.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False));raise SystemExit(0 if report['status'].startswith('PASS') else 1)
